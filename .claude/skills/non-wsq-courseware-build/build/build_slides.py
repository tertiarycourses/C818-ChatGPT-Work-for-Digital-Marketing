#!/usr/bin/env python3
"""C818 visual engine. Uses the supplied reference theme; content from course.json/labs.json.
No detailed lab procedures appear in the deck. Every concept precedes its topic labs.
"""
from pathlib import Path
import json,math
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
import course_data as C
R=C.ROOT; D=C.DATA;LABS=C.LABS
P=Presentation();P.slide_width=Inches(13.333333);P.slide_height=Inches(7.5)
INK='21201C';CREAM='FAF9F5';ACC='D97757';TEAL='3A7D7B';PANEL='2E2C26';GREY='67645D';IVORY='F3EFE7';WHITE='FFFFFF'
log=[]
def rgb(x):return RGBColor.from_string(x)
def rect(s,x,y,w,h,color,kind=1):
 sp=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));sp.fill.solid();sp.fill.fore_color.rgb=rgb(color);sp.line.fill.background();return sp
def text(s,x,y,w,h,txt,size=20,color=INK,bold=False,font='Calibri',align=None):
 tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(str(txt).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(6)
  if align:p.alignment=align
  r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
 return tb
def slide(title,tag='CONCEPT',dark=False,typ='concept'):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(INK if dark else CREAM)
 text(s,.65,.42,11.6,.35,tag,12,ACC,True);text(s,.65,.96,12,.9,title,30,IVORY if dark else INK,True,'Cambria')
 text(s,.65,7.12,7,.2,'© Tertiary Infotech Academy Pte Ltd',8,'A39E92' if dark else GREY)
 text(s,7.1,7.12,5.2,.2,'C818 · ChatGPT WORK for Digital Marketing',8,'A39E92' if dark else GREY,align=PP_ALIGN.RIGHT)
 text(s,12.5,7.12,.3,.2,str(len(P.slides)),8,'A39E92' if dark else GREY)
 # Restrained fade; static PDF contains complete meaning.
 tr=OxmlElement('p:transition');tr.set('spd','med');tr.append(OxmlElement('p:fade'));s._element.append(tr)
 log.append(dict(num=len(P.slides),title=title,type=typ,tag=tag))
 return s
def notes(s,msg):s.notes_slide.notes_text_frame.text=msg

def callout(s,txt,dark=False):
 rect(s,.65,6.35,12.02,.48,PANEL if dark else 'ECE8E0',5)
 text(s,.85,6.46,11.6,.25,txt,11,IVORY if dark else INK)

def cards(s,items,y=2.8,dark=False):
 n=len(items);gap=.25;w=(12.03-gap*(n-1))/n
 for i,(head,body) in enumerate(items):
  x=.65+i*(w+gap);rect(s,x,y,w,2.45,PANEL if dark else WHITE,5)
  rect(s,x+.18,y+.22,.45,.45,ACC,9);text(s,x+.18,y+.25,.45,.35,str(i+1),15,IVORY,True,align=PP_ALIGN.CENTER)
  text(s,x+.82,y+.28,w-1.0,.6,head,21,IVORY if dark else INK,True,'Cambria')
  text(s,x+.25,y+1.02,w-.5,1.15,body,18,IVORY if dark else GREY)

def flow(s,labels):
 w=3.65
 for i,it in enumerate(labels):
  head,body=it.split('|',1);x=.65+i*4.16
  rect(s,x,2.9,w,1.85,WHITE,5);rect(s,x+.18,3.08,.55,.55,ACC,9)
  text(s,x+.18,3.17,.55,.35,str(i+1),18,IVORY,True,align=PP_ALIGN.CENTER)
  text(s,x+.85,3.15,w-1.05,.5,head,23,INK,True,'Cambria');text(s,x+.24,3.95,w-.48,.6,body,19,GREY)
  if i<2:rect(s,x+w+.1,3.6,.24,.3,TEAL,33)

def ledger(s,labels,rows=None):
 heads=['Item','Evidence / criterion','Decision'];rows=rows or [[x.split('|')[0],x.split('|')[1],['Use','Propose','Check'][i]] for i,x in enumerate(labels)]
 widths=[2.3,6.3,3.4];xs=[.65,2.95,9.25]
 for j,h in enumerate(heads):rect(s,xs[j],2.95,widths[j],.55,INK);text(s,xs[j]+.15,3.07,widths[j]-.3,.3,h,15,IVORY,True)
 for i,row in enumerate(rows):
  for j,val in enumerate(row):rect(s,xs[j],3.53+i*.66,widths[j],.62,WHITE if i%2==0 else 'ECE8E0');text(s,xs[j]+.15,3.67+i*.66,widths[j]-.3,.4,val,16,INK)

def picture(s,name,x,y,w,h):
 p=R/'courseware/assets'/name
 if not p.exists():return
 from PIL import Image
 iw,ih=Image.open(p).size;scale=min(w/iw,h/ih);ww,hh=iw*scale,ih*scale
 s.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def special(s,c):
 layout=c['layout']; labels=c['labels']
 if layout=='flow':flow(s,labels)
 elif layout in ['ledger','calendar']:
  if layout=='calendar':ledger(s,labels,[['Day 1','Teach: useful learning question','Draft'],['Day 3','Prepare: what to bring','Draft'],['Day 9','Invite: source-checked course offer','Draft']])
  else:ledger(s,labels)
 elif layout=='palette':
  for i,(name,col) in enumerate([('Crust','B5652D'),('Ink','2B211B'),('Cream','FBF7F1'),('Basil','3F6B4A')]):
   rect(s,.65+i*3.05,3.0,2.8,1.4,col,5);text(s,.65+i*3.05,4.65,2.8,.8,name+'\n#'+col,20,INK,True,'Cambria')
 elif layout in ['site','email','flyer']:
  pic={'site':'academy-desktop.png','email':'email-desktop.png','flyer':'flyer-a5.png'}[layout]
  picture(s,pic,7.65,2.55,4.7,3.5)
  for i,l in enumerate(labels):
   h,b=l.split('|');rect(s,.65,2.8+i*1.05,6.4,.85,WHITE,5);text(s,.85,2.93,2,.45,h,18,TEAL,True,'Cambria');text(s,3,2.96,3.7,.45,b,17,INK)
 elif layout=='metrics':
  for i,(label,val) in enumerate([('Variant A',5.0),('Variant B',5.95)]):
   y=3.25+i*1.25;text(s,.65,y,2.3,.5,label,23,INK,True,'Cambria');rect(s,3,y,val*1.1,.55,ACC if i==0 else TEAL);text(s,3+val*1.1+.15,y,.95,.5,f'{val:g}%',23,INK,True)
  text(s,.65,5.85,11.8,.35,'Unique clicks ÷ delivered emails · Synthetic teaching figures',14,GREY)
 else:cards(s,[tuple(x.split('|',1)) for x in labels])

# Cover mirrors the supplied reference's dark cover, serif heading and corner discs.
s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(INK)
rect(s,10.75,-1.5,4.0,4.0,'3F2D24',9);rect(s,11.6,4.9,3.2,3.2,'233530',9)
text(s,.8,1.05,9.4,.5,'NON-WSQ COURSE · C818',14,ACC,True)
text(s,.8,1.95,11.3,2.1,'ChatGPT WORK\nfor Digital Marketing',43,IVORY,True,'Cambria')
text(s,.8,4.45,10,.7,'Flyer design · Social articles · Email campaigns',22,IVORY)
text(s,.8,5.5,10,.6,'Tertiary Infotech Academy Pte Ltd\nUEN: 201200696W',13,IVORY)
text(s,.8,6.42,11,.35,'Version 1.0 · 1 October 2026 · 2 days / 15 instructional hours',12,'A39E92')
log.append(dict(num=1,title=D['title'],type='admin',tag='COVER'));notes(s,'C818. The entire course uses the existing fictional Cook & Bake Academy catalogue. Classroom drafts only. Theme follows the supplied Agentic AI reference deck. Detailed steps are in the LG and lab folders.')
s=slide('What you will build','COURSE OVERVIEW');cards(s,[('Flyer package','An editable A5 flyer, source proof and square/story derivatives.'),('Social package','A developed article, distinct captions and five-post calendar.'),('Email package','Three unsent emails, eligible synthetic audience and metric handoff.')]);callout(s,'One shared campaign · the same catalogue facts across every channel')
s=slide('Learning outcomes','COURSE OVERVIEW');
for i,lo in enumerate(C.LEARNING_OUTCOMES):
 x=.65+(i%2)*6.15;y=2.1+(i//2)*1.85;rect(s,x,y,5.88,1.6,WHITE,5);text(s,x+.22,y+.22,5.42,1.22,lo,20,INK)
notes(s,'Ask learners which deliverable is most useful to their current role. Map LO1 to Labs 1–2, LO2 to Labs 3–4, LO3 to Labs 5–6 and LO4 to Labs 7–8.')
s=slide('How You’ll Learn','LEARNING FLOW');flow(s,['Demonstrate|Concept and worked example','Build + verify|Create the asset and prove it','Discuss + recap|Review a choice and improve it']);callout(s,'Detailed instructions and prompts live in the Learner Guide and lab READMEs.')
s=slide('Two days, one connected campaign','COURSE ROUTE');cards(s,[('Day 1 · Ground','Labs 1–2: source facts, brief and reusable workflow.'),('Day 1 · Design','Labs 3–4: A5 flyer, variants and visible-text proof.'),('Day 2 · Connect','Labs 5–8: social, email, synthetic metrics and handoff.')]);callout(s,'Daily plan: 09:30–18:30 · 60-minute lunch · 30 minutes of tea breaks · 450 instructional minutes')
s=slide('Bring the right tools and context','SETUP');cards(s,[('Work access','A ChatGPT account with Work available; features vary by workspace.'),('Preview tools','Browser and layout editor. Python 3 previews the supplied website.'),('Recovery route','Each lab folder includes source files, prompts and a supplied checkpoint.')]);callout(s,'Use project Sources on web or an approved local project folder on desktop.')
s=slide('Ground rules for useful work','HOUSEKEEPING');cards(s,[('Participate','Ask questions, compare choices and return on time.'),('Use synthetic data','Only fictional catalogue and example.com contact records.'),('Keep work reviewable','Save editable files and evidence; all campaigns stay unsent drafts.')]);callout(s,'Unknown offers stay UNKNOWN. Review each saved file before using it.')
s=slide('The existing Cook & Bake Academy site','SHARED SCENARIO');picture(s,'academy-desktop.png',.65,2.1,8.1,4.7);text(s,9.1,2.5,3.55,3.3,'20 synthetic courses\n10 Bakery + 10 Cooking\n\nCampaign spine:\nBAK-101 Sourdough\n\nCafe-business angle is a proposal.',20,INK);notes(s,'The C818 copy preserves the original catalogue, brand and signup demo. Optional assistant is adapted to source-only search, without API credential entry. Source files are in Lab 1 assets/cnbacademy.')
for t in D['topics']:
 s=slide(t['title'],f"TOPIC {t['num']} · {t['lo']}",True,'section');rect(s,10.5,4.8,4,4,'3F2D24',9);text(s,.65,2.5,10.8,1.1,t['subtitle'],25,IVORY);text(s,.65,4.45,10.8,1.0,'Understand the model. Apply it to the academy. Verify the saved result.',25,IVORY,font='Cambria');notes(s,'Teach the concepts before introducing the labs. Learners retain the same selected course and ledger throughout.')
 for c in t['concepts']:
  s=slide(c['title'],f"TOPIC {t['num']} · CONCEPT")
  # Short teaching statement; deeper explanations remain in notes and LG.
  first=c['explanation'].split('. ')[0]+'.';text(s,.65,1.98,12.0,.62,first,18,INK)
  special(s,c);callout(s,'CHECK · '+c['pitfall'])
  notes(s,c['explanation']+'\n\nWorked example: '+c['example']+'\n\nCommon mistake: '+c['pitfall'])
  # Worked example strip is concise and substantive, without duplicating the prose.
  if c['layout'] not in ['metrics','palette','site','email','flyer']:
   ex=c['example'];text(s,.65,5.53,12.0,.62,ex,15,GREY)
 for a in [x for x in LABS if x['topic']==t['num']]:
  s=slide(f"Lab {a['num']:02d} · {a['title']}",f"LAB OVERVIEW · {a['duration']} MINUTES",True,'lab')
  text(s,.65,2.0,11.8,.65,a['goal'],23,IVORY)
  cards(s,[('What you’ll build',a['build'].split('. ')[0]),('Why it matters','Apply '+t['title'].lower()+' to the same source-checked campaign.'),('Finished evidence','Open the files and compare the required facts and checks in the lab guide.')],y=3.0,dark=True)
  callout(s,f"Detailed procedure: labs/lab-{a['num']:02d}-{a['slug']}/README.md",True)
  notes(s,a['build']+'\nExpected checks: '+' '.join(a['checks'])+'\nUse the supplied checkpoint if a learner falls behind. The PPT intentionally contains no detailed procedures.')
 s=slide(f"Topic {t['num']} recap",f"RECAP · {t['lo']}",False,'recap');cards(s,[('Explain',t['concepts'][0]['title']+'. Explain why it matters to the campaign.'),('Show','Open the current lab outputs and identify the evidence that supports their facts.'),('Improve','Name one correction and the next check another marketer can reproduce.')]);callout(s,C.LEARNING_OUTCOMES[t['num']-1]);notes(s,'Invite learners to explain their choices using the actual artifacts rather than an agent summary.')
s=slide('What you achieved','WRAP-UP',False,'closing');cards(s,[('Consistent offer','The flyer, social article and email all use one selected catalogue record.'),('Reviewable assets','Editable sources, proofs, audience rule and manifest are saved.'),('Evidence-based next step','Synthetic metrics support a proposed test, with limitations stated.')]);callout(s,'Completing the classroom package is distinct from launching a real campaign.')
s=slide('Keep practising','NEXT STEPS',False,'closing');flow(s,['Repeat|Choose a second catalogue course','Compare|Re-run the claim and format checks','Apply|Use an authorised real brief']);callout(s,'Confirm current facts, tools, audience and distribution approval before production use.')
s=slide('Sources and further reading','LEARNING RESOURCES',False,'closing');ledger(s,[],[['OpenAI Learn','ChatGPT Work, projects, skills and image guidance','1 Oct 2026'],['W3C WAI','Contrast minimum guidance','1 Oct 2026'],['Mailchimp','Open and click rate definitions and limitations','1 Oct 2026']]);text(s,.65,5.8,12,.3,'Clickable source URLs are recorded in the Learner Guide and SOURCE-NOTES.md.',15,GREY);notes(s,'Official sources: https://learn.chatgpt.com/docs/get-started-with-work ; https://learn.chatgpt.com/docs/projects ; https://learn.chatgpt.com/docs/skills-and-plugins ; https://learn.chatgpt.com/docs/image-generation ; https://www.w3.org/WAI/WCAG21/Understanding/contrast-minimum ; https://mailchimp.com/help/about-open-and-click-rates/')
s=slide('Thank you','C818 · TERTIARY INFOTECH',True,'closing');rect(s,10.7,4.7,4,4,'233530',9);text(s,.65,2.6,11.4,1.5,'Make useful work.\nKeep the evidence with it.',36,IVORY,True,'Cambria');text(s,.65,5.2,11.4,.6,'www.tertiarycourses.com.sg · enquiry@tertiaryinfotech.com',18,IVORY)
out=R/'courseware'/f'{C.SHORT_TITLE}-{C.VERSION}.pptx';P.save(out);(R/'data/slide-manifest.json').write_text(json.dumps(log,indent=2));print('Saved',out,'slides',len(P.slides))

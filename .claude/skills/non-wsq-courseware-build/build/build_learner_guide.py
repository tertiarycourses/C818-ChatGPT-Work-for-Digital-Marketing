#!/usr/bin/env python3
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import course_data as C
import prodoc
R=C.ROOT;D=C.DATA;L=C.LABS
B=[]
def add(k,*v):B.append((k,*v))
add('h1','Introduction')
add('p','This guide develops the concepts behind a source-grounded digital marketing workflow and provides complete procedures for eight connected labs. You will create an editable flyer package, social article and content plan, an unsent email campaign and a reviewable handoff. The fictional Cook & Bake Academy website and its catalogue are the shared scenario throughout.')
add('p','The course is two days with 15 instructional hours. The lesson plan uses 450 instructional minutes plus 30 minutes of tea breaks each day, excluding a 60-minute lunch. Work features depend on account and workspace access. A file-upload route and supplied recovery packs keep the labs usable when local tools or image generation are unavailable.')
add('h2','Learning outcomes')
for x in C.LEARNING_OUTCOMES:add('bullet',x)
add('h2','Tools, access and working folders')
add('p','Use a ChatGPT account with Work available, a browser, a text editor and a layout editor such as PowerPoint or Google Slides. Python 3 serves the existing teaching website locally. No real subscriber list, connected email account, paid advertising account or API credential is required.')
add('code','C818-marketing/\n  cnbacademy/            # copy from Lab 1 assets\n  campaign/              # brief, ledger, metrics, manifest\n  creative/sourdough/    # editable flyers and proofs\n  content/               # article, post packages, calendar\n  email/                 # synthetic segment and unsent sequence')
add('p','On web, add the relevant sources to the project’s Sources section and download the generated files. On desktop, an approved local project can read and write the attached working folder. Ask Work to list the sources it can read before starting. Merely pasting a path does not grant access.')
add('h2','The reference workflow translated to marketing')
add('table',['Stage','Marketing meaning','Evidence'],[['Goal','Audience, offer and one action','brief.md'],['Plan','Compare routes before production','design-decision.md'],['Create','Make the requested editable assets','Source files and proofs'],['Preserve context','Keep brand and source facts available','workflow.md and ledger'],['Version','Identify current deliverables','manifest.csv'],['Prepare distribution','Resolve production settings separately','send-checklist.md'],['Verify','Open and compare saved artifacts','Review logs']])
add('h2','The shared academy scenario')
add('p','The existing Cook & Bake Academy site has 20 synthetic courses: 10 Bakery and 10 Cooking. BAK-101 Artisan Sourdough Bread Baking is the sample campaign spine. Its four-week duration, S$680 fee and Saturday schedule are teaching facts. A cafe-business audience angle may be proposed, but no cafe/barista course is verified in this catalogue. Unknown discounts, remaining seats, tax treatment and public booking links must stay UNKNOWN.')
add('image','academy-desktop.png','Existing academy website, captured from the C818 teaching copy. Not a real public offer.')
for t in D['topics']:
 add('h1',f"Topic {t['num']} — {t['title']}");add('p',t['subtitle']+'. '+C.LEARNING_OUTCOMES[t['num']-1])
 for c in t['concepts']:
  add('h2',c['title']);add('p',c['explanation']);add('table',['Element','Meaning'],[x.split('|',1) for x in c['labels']]);add('p','Worked example: '+c['example']);add('p','Common mistake: '+c['pitfall'])
 # Actual visual evidence, plus table-based frameworks throughout the topic.
 if t['num']==2:add('image','flyer-a5.png','Supplied editable A5 classroom sample; inspect the typography and source facts.')
 if t['num']==4:add('image','email-desktop.png','Supplied HTML email preview; a browser preview is not an inbox deliverability test.')
 for a in [x for x in L if x['topic']==t['num']]:
  add('h2',f"Lab {a['num']:02d} — {a['title']}");add('p',f"Duration: {a['duration']} minutes. Learning outcome: {t['lo']}.");add('p','Goal: '+a['goal']);add('p','What you will build: '+a['build']);add('p','Prerequisites: '+a['prerequisites'])
  add('h3','Detailed procedure')
  for i,(h,txt) in enumerate(a['steps'],1):add('p',f'{i}. {h}. {txt}')
  add('h3','Copy-ready task prompt');add('code',a['prompt'])
  add('h3','Test it')
  for check in a['checks']:add('bullet',check)
  add('h3','Troubleshooting')
  for issue,fix in a['troubleshooting']:add('p',issue+': '+fix)
  add('h3','Challenge');add('p',a['challenge']);add('h3','Reflection');add('p',a['reflection'])
  path=f"labs/lab-{a['num']:02d}-{a['slug']}";add('p','Recovery checkpoint: '+path+'/checkpoints/. Supplied files are sample classroom drafts. Record sample provenance and verify the facts before continuing. Each lab has its own README, assets and prompts in '+path+'/ .')
 add('h2',f"Topic {t['num']} recap — {t['lo']}");add('p','Explain the first concept in your own words, open the current lab output and show the evidence for its factual claims. Name one revision and the next check another marketer can reproduce. '+C.LEARNING_OUTCOMES[t['num']-1])
add('h1','Worked data and reusable templates')
selected=json.loads((R/'labs/lab-01-grounded-campaign-brief/checkpoints/selected-course.json').read_text())
add('h2','Selected source record — BAK-101');add('code',json.dumps(selected,indent=2,ensure_ascii=False));add('p','All values above are copied from the fictional teaching catalogue, snapshot 1 October 2026. They are not Tertiary Infotech course offers.')
add('h2','Claim ledger example');add('table',['Claim','Value','Source','Status'],[['Fee','S$680','BAK-101.fee','Scenario fact'],['Duration','4 weeks','BAK-101.weeks','Scenario fact'],['Schedule','Saturday 9:30am–12:30pm','BAK-101.when','Scenario fact'],['Seats remaining','UNKNOWN','Not supplied','Hold'],['Public booking URL','UNKNOWN','Not supplied','Hold']])
add('h2','Synthetic contacts and eligibility')
add('code',(R/'labs/lab-07-email-sequence-and-segmentation/assets/contacts.csv').read_text());add('p','Include only consent=yes AND interest=bakery AND suppressed=no. The exact included IDs are C001, C004 and C006. Missing consent is excluded. No real addresses or credentials are needed.')
add('h2','Synthetic results and denominator check');add('table',['Variant','Sent','Delivered','Unique clicks','Bookings'],[['A','1000','980','49','7'],['B','1000','975','58','8']]);add('p','A: 980/1000 = 98.0% delivery; 49/980 = 5.0% unique click rate; 7/49 ≈ 14.3% click-to-booking. B: 975/1000 = 97.5%; 58/975 ≈ 5.95%; 8/58 ≈ 13.79%. A zero denominator produces UNKNOWN. These figures are invented teaching data; they do not establish real academy performance or statistical confidence.')
add('h2','Review and handoff templates');add('code','Review: file,criterion,evidence,status,correction\nCalendar: day,platform,pillar,audience,asset,source,cta,owner,dependency,status,hard_sell\nManifest: path,purpose,version,source,status,unresolved_issue\nHold: unknown,impact,owner,next_action')
add('h1','Keep practising')
add('p','Repeat the connected campaign for a second catalogue course, preserving a separate source ledger. Compare the visual and editorial choices against a different audience need. Before applying the workflow to a real campaign, confirm the current facts, approved tools, recipient eligibility and distribution settings with the campaign owner.')
add('h2','Glossary');add('table',['Term','Meaning'],[['Claim ledger','A field-level record of factual evidence, proposals and unknowns.'],['Hierarchy','The visual reading order established by size, position and contrast.'],['Repurposing','Adapting one useful idea to a different format while preserving meaning.'],['Suppression','A rule that excludes a contact from the eligible audience.'],['Unique click rate','For this exercise, unique clickers divided by delivered emails.'],['Checkpoint','A supplied recovery pack that allows a learner to rejoin.'],['Handoff manifest','An inventory of current files, evidence, status and unresolved issues.']])
add('h1','Sources and further reading')
refs=[('Official C818 course page','https://www.tertiarycourses.com.sg/chatgpt-work-for-digital-marketing.html'),('OpenAI: Get started with ChatGPT Work','https://learn.chatgpt.com/docs/get-started-with-work'),('OpenAI: Projects and chats','https://learn.chatgpt.com/docs/projects'),('OpenAI: Skills and plugins','https://learn.chatgpt.com/docs/skills-and-plugins'),('OpenAI: Image generation','https://learn.chatgpt.com/docs/image-generation'),('W3C WAI: Contrast minimum','https://www.w3.org/WAI/WCAG21/Understanding/contrast-minimum'),('Mailchimp: Open and click rates','https://mailchimp.com/help/about-open-and-click-rates/')]
add('p','Checked 1 October 2026. The official registration page establishes C818, the title, two days and 15 hours. Its topic text currently describes Codex engineering; this package follows the user-requested ChatGPT Work marketing scope. Product features remain dependent on account, platform and workspace. The flyer formats, caption limits, social mix and synthetic metrics are classroom brief choices, not universal platform rules.')
for name,url in refs:add('p',name+': '+url)
# Markdown and Word share the exact content stream.
md=[f'# {C.TITLE} — Learner Guide',f'C818 · Version 1.0 · {C.VERSION_DATE}']
for k,*v in B:
 if k in ['h1','h2','h3']:md.append('#'*(int(k[1])+1)+' '+v[0])
 elif k=='p':md.append(v[0])
 elif k=='bullet':md.append('- '+v[0])
 elif k=='code':md.append('```text\n'+v[0].strip()+'\n```')
 elif k=='table':md.append('| '+' | '.join(v[0])+' |\n| '+' | '.join(['---']*len(v[0]))+' |\n'+'\n'.join('| '+' | '.join(str(x) for x in row)+' |' for row in v[1]))
 elif k=='image':md.append(f'![{v[1]}](courseware/assets/{v[0]})')
(R/f'LG-{C.SHORT_TITLE}.md').write_text('\n\n'.join(md)+'\n')
doc=Document();normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
prodoc.style_headings(doc);prodoc.add_cover_page(doc,'LEARNER GUIDE',C.TITLE,'1.0',org_logo=str(Path(__file__).parents[1]/'assets/tertiary-infotech-logo.png'),course_code=C.COURSE_CODE);prodoc.add_version_control(doc,C.VERSION_HISTORY);prodoc.add_toc(doc)
for k,*v in B:
 if k in ['h1','h2','h3']:doc.add_heading(v[0],level=int(k[1]))
 elif k=='p':doc.add_paragraph(v[0])
 elif k=='bullet':doc.add_paragraph(v[0],style='List Bullet')
 elif k=='code':
  for line in v[0].strip().splitlines():
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.left_indent=Pt(8);r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1EDE5');p._p.get_or_add_pPr().append(shade)
 elif k=='table':
  table=doc.add_table(rows=1,cols=len(v[0]));table.style='Table Grid'
  for cell,h in zip(table.rows[0].cells,v[0]):cell.text=h;prodoc._shade_cell(cell,'F1EDE5')
  for row in v[1]:
   for cell,value in zip(table.add_row().cells,row):cell.text=str(value)
  for row in table.rows:
   for cell in row.cells:
    for p in cell.paragraphs:
     for r in p.runs:r.font.size=Pt(10)
 elif k=='image':
  p=R/'courseware/assets'/v[0]
  if p.exists():doc.add_picture(str(p),width=Inches(5.7));cap=doc.add_paragraph(v[1]);cap.style='Caption'
prodoc.add_page_numbers(doc);prodoc.enable_update_fields(doc);out=R/'courseware'/f'LG-{C.SHORT_TITLE}.docx';doc.save(out);print('Saved LG',out)

from pathlib import Path
import json,csv,re
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[1];L=json.loads((R/'data/labs.json').read_text()); D=json.loads((R/'data/course.json').read_text())
def folder(n):
 a=L[n-1];return R/'labs'/f"lab-{n:02d}-{a['slug']}"
def write(n,path,text):
 p=folder(n)/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text)
site=folder(1)/'assets/cnbacademy';courses=json.loads((site/'data/courses.json').read_text());selected=next(c for c in courses if c['code']=='BAK-101')
# Retain original catalogue, visual theme and signup; simplify optional assistant to source-only lookup.
s=(site/'index.html').read_text();s=re.sub(r'<button class="chat-settings-toggle".*?</button>','',s);s=re.sub(r'<form class="chat-settings".*?</form>','',s,flags=re.S)
s=s.replace('<body>','<body>\n<div style="padding:8px;text-align:center;background:#2B211B;color:#FBF7F1;font:14px sans-serif">C818 fictional teaching website · Demo data only</div>')
(site/'index.html').write_text(s)
(site/'js/chat.js').write_text('''const panel=document.querySelector('#chat-panel'),log=document.querySelector('#chat-log'),form=document.querySelector('#chat-form'),input=document.querySelector('#chat-input');
let courses=[];fetch('data/courses.json').then(r=>r.json()).then(d=>courses=d);
function message(text){const p=document.createElement('p');p.className='chat-message';p.textContent=text;log.append(p);}
document.querySelectorAll('[data-open-chat]').forEach(b=>b.addEventListener('click',()=>{panel.hidden=false;input.focus();if(!log.children.length)message('Search the fictional catalogue by course code or title. No API key is needed.');}));
document.querySelector('[data-close-chat]').addEventListener('click',()=>panel.hidden=true);
form.addEventListener('submit',e=>{e.preventDefault();const q=input.value.trim().toLowerCase();message(input.value);input.value='';const hits=courses.filter(c=>q.includes(c.code.toLowerCase())||c.title.toLowerCase().split(' ').some(w=>w.length>4&&q.includes(w)));message(hits.length?hits.slice(0,3).map(c=>`${c.code}: ${c.title}. S$${c.fee}; ${c.weeks} week(s); ${c.when}; ${c.campus}. Source: data/courses.json (fictional teaching data).`).join('\\n'):'UNKNOWN in the supplied catalogue. Search a course code such as BAK-101.');});
''')
(site/'TEACHING-NOTE.md').write_text((site/'TEACHING-NOTE.md').read_text()+'\nThe C818 copy preserves the catalogue and visual system. Its optional assistant has been adapted to catalogue-only lookup, with all API credential entry and network model calls removed.\n')
# remove unused original API RAG module from learner copy
(site/'js/rag.js').unlink(missing_ok=True)
write(1,'assets/project-instructions.md','''# C818 project guidance

Use ChatGPT Work to produce reviewable marketing files. Read the academy brand, brief and catalogue before drafting. Use only the selected course record for offer facts. Separate source facts, creative proposals and UNKNOWNs. Keep the exact tagline Learn it. Make it. Take it further. Do not invent discounts, remaining seats, trainer identities, public booking links or a cafe course. Preserve editable factual text and return source checks. Use synthetic data only. Plan before production and stop before posting, scheduling or sending. Keep outputs as classroom drafts.
''')
brief='''# BAK-101 campaign brief — supplied classroom draft

Audience: weekend home bakers who want a reliable starter routine. Secondary proposed angle: bakery skills for someone exploring a cafe business; no cafe course is claimed.
Need: confidence maintaining a starter and practising bread techniques.
Objective: encourage readers to view the course details; numeric target UNKNOWN.
Core message: Learn to feed and maintain a starter and bake country loaves through structured practice.
Offer: BAK-101 Artisan Sourdough Bread Baking; 4 weeks; S$680; Saturday 9:30am–12:30pm; Orchard Road. These are fictional catalogue facts.
CTA: View course details in the demo catalogue.
Scope: A5 flyer plus square/story variants, social article and five-post plan, three-email draft sequence.
Distribution: classroom draft only; no live posting, scheduling or send.
'''
write(1,'checkpoints/brief.md',brief);write(1,'checkpoints/selected-course.json',json.dumps(selected,indent=2))
ledger='claim,value,source_field,status\n'
for k in ['code','title','fee','weeks','when','campus','allergens']:ledger+=f'"{k}","{selected[k]}","data/courses.json BAK-101.{k}","scenario fact"\n'
for k in ['remaining seats','discount','tax treatment','public booking link','sales target']:ledger+=f'"{k}","UNKNOWN","absent","UNKNOWN"\n'
write(1,'checkpoints/claim-ledger.csv',ledger);write(1,'checkpoints/unknowns.md','# Unknowns — classroom draft\n\nRemaining seats, discounts, tax treatment, public booking link, campaign performance and sales target are UNKNOWN. Confirm with the authorised campaign owner before any production use.\n')
workflow='''# Marketing production workflow — supplied classroom draft

1. Goal: input the catalogue and audience need; output a brief; check one offer and one action.
2. Plan: input the brief; output two routes; check both preserve facts before choosing.
3. Create: input the selected route; output editable assets; check requested formats exist.
4. Preserve context: input brand and source ledger; output project guidance; check drafts never become offer sources.
5. Version: input revised files; output a current-file manifest; check the approved copy is identifiable.
6. Prepare distribution: input current assets; output an unsent handoff; check sender, audience and live destination are separately reviewed.
7. Verify: input artifact and sources; output proof and review log; check every visible fact and unresolved UNKNOWN.

Academy tagline: Learn it. Make it. Take it further.
Voice: warm, practical, confident. Palette: Crust #B5652D, Ink #2B211B, Cream #FBF7F1, Basil #3F6B4A. Headings: Georgia; body: sans-serif.
Facts: cnbacademy/data/courses.json only. No term highlights have been supplied.
Reuse: add this file to project Sources and explicitly refer to it. Uploading a skill example does not install it.
All outputs remain classroom drafts; do not post, schedule or send.
'''
write(2,'checkpoints/workflow.md',workflow);write(2,'checkpoints/brand-check.md','# Brand check\n\nCheck exact tagline, practical wording, academy palette, source facts, editable copy, actual-image alt text and absence of unsupported urgency. Record issue, correction and evidence.\n')
write(2,'checkpoints/workflow-trial.md','# Sample trial\n\nBuild a practical starter routine with Artisan Sourdough Bread Baking. In the fictional BAK-101 catalogue, the four-week class covers starter care, shaping, proofing and home-oven baking. View course details in the demo catalogue. This is a classroom draft, not a live offer.\n\nSource: BAK-101.title, weeks and learn.\n')
copy='''Cook & Bake Academy
Build your sourdough routine
Artisan Sourdough Bread Baking
BAK-101 · 4 weeks · S$680
Saturday 9:30am–12:30pm
Orchard Road
Practise starter care, shaping and home-oven baking.
View course details in the demo catalogue
Learn it. Make it. Take it further.
Classroom draft · Fictional teaching offer
'''
write(3,'assets/design-brief.md','# Flyer brief\n\nAudience: weekend home baker. Format: A5 portrait, 148 × 210 mm. Safe area: 10 mm. One action: view details in the demo catalogue. Use source-checked BAK-101 copy and the academy brand. Image optional; all factual text editable. Printer bleed, live booking link and tax treatment UNKNOWN.\n')
write(3,'checkpoints/flyer-copy.md',copy);write(3,'checkpoints/image-prompt.txt','Optional illustration prompt: Natural-light sourdough loaf on a cream work surface, hands out of frame, generous clear space, warm practical food photography. No lettering, faces, logos or implied real class evidence.\nNo image was generated for the supplied typographic sample.\n')
write(3,'checkpoints/text-proof.md','# Visible text proof — sample\n\n'+copy+'\nSources: title/code/weeks/fee/when/campus/learn from BAK-101. Tagline from the brand-voice example. Headline and CTA are creative proposals.\n')
def rgb(x):return RGBColor.from_string(x)
def box(s,x,y,w,h,text,size=18,bold=False,color='2B211B',font='Arial'):
 tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_top=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(5);r=p.add_run();r.text=line;r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
 return tb
for n,fmt,w,h in [(3,'a5',5.8268,8.2677),(4,'square',7.874,7.874),(4,'story',7.0866,12.5984)]:
 p=Presentation();p.slide_width=Inches(w);p.slide_height=Inches(h);s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('FBF7F1')
 bar=s.shapes.add_shape(1,0,0,Inches(w),Inches(.16));bar.fill.solid();bar.fill.fore_color.rgb=rgb('B5652D');bar.line.fill.background()
 x=.5;area=w-1;sc=w/5.8268
 box(s,x,1.35 if fmt=='story' else .55,area,.45,'COOK & BAKE ACADEMY',14*sc,True,'3F6B4A')
 box(s,x,2.05 if fmt=='story' else 1.22,area,1.7,'Build your\nsourdough routine',28*sc,True,font='Georgia')
 # editorial typographic motif, fully editable
 shape=s.shapes.add_shape(1,Inches(x),Inches(4.1 if fmt=='story' else 3.1),Inches(area),Inches(1.0 if fmt!='story' else 2.0));shape.fill.solid();shape.fill.fore_color.rgb=rgb('3F6B4A');shape.line.fill.background()
 box(s,x+.22,4.4 if fmt=='story' else 3.3,area-.44,1.3,'STARTER  /  SHAPE  /  BAKE',13*sc,True,'FBF7F1')
 y=6.65 if fmt=='story' else (4.15 if fmt=='square' else 4.35)
 box(s,x,y,area,.8,'Artisan Sourdough\nBread Baking',19*sc,True,font='Georgia')
 box(s,x,y+1.0,area,1.2,'BAK-101 · 4 weeks · S$680\nSaturday 9:30am–12:30pm\nOrchard Road',13*sc)
 box(s,x,9.8 if fmt=='story' else (6.4 if fmt=='square' else h-1.62),area,.5,'View course details in the demo catalogue',12*sc,True,'3F6B4A')
 box(s,x,10.55 if fmt=='story' else h-1.02,area,.3,'Learn it. Make it. Take it further.',11*sc,False)
 box(s,x,11.0 if fmt=='story' else h-.55,area,.25,'Classroom draft · Fictional teaching offer',9*sc)
 path=folder(n)/'checkpoints'/f'flyer-{fmt}.pptx';p.save(path)
write(4,'assets/design-review-template.csv','file,criterion,evidence,status,correction\nflyer-square.pdf,source facts,,draft,\nflyer-square.pdf,one CTA,,draft,\nflyer-square.pdf,readability,,draft,\nflyer-story.pdf,source facts,,draft,\nflyer-story.pdf,safe area,,draft,\nflyer-story.pdf,contrast,,draft,\n')
write(4,'checkpoints/variants-proof.md','# Variant proof — supplied sample\n\nSquare and story contain the same factual copy as the A5 sample. Inspect each PDF independently before use. The layouts are typographic, not photographic, and have no working booking code.\n')
article='''# Building confidence with a sourdough starter routine

## Start with the skill you want to practise

A sourdough starter can make baking feel like an ongoing project rather than a single recipe. If you are deciding whether to join a class, begin with the skill you want to practise. Is your main question how to maintain a routine, how to recognise a dough stage, or how to shape a loaf with more confidence? Naming that question helps you choose a learning experience for a clear reason.

A course description should help you make that choice. Look for concrete learning outcomes rather than a promise that every loaf will look the same. A useful outcome describes an action you can practise and discuss. It also gives you something to return to when you review what you have learned. Confidence is a proposed benefit of practice, not a guarantee that a course can make for every person.

For a home baker, the goal may be to understand a repeatable process well enough to ask better questions. For someone exploring a cafe business, baking skills may be one part of a larger plan. That second angle does not turn a bakery course into a cafe qualification. Keep the learning need and the actual course offer separate.

## Look for a connected learning process

The fictional Cook & Bake Academy catalogue describes BAK-101 Artisan Sourdough Bread Baking through a set of practical learning outcomes. These include feeding and maintaining a wild-yeast starter, mixing and shaping by hand, reading proof through practical checks, and scoring and baking in a home oven with a Dutch oven. Those outcomes describe what the teaching scenario covers; they are not a recipe in this article.

The connection between those outcomes matters. A learner can think about the starter, dough handling and baking as related parts of one process instead of collecting isolated tips. When reading a course page, ask whether its learning outcomes address your own question and whether the level and schedule fit your situation. A longer list of techniques is not automatically a better fit.

It is also useful to notice what the description does not tell you. The catalogue does not provide a promise of professional certification, guaranteed business income or a claim that places are about to sell out. Those details should not appear in a marketing draft simply because they might sound persuasive. Missing information is a reason to ask a question, not a reason to invent an answer.

## Choose with the details in view

In this classroom scenario, BAK-101 is a beginner course lasting four weeks, with a fee of S$680. Its listed schedule is Saturday 9:30am–12:30pm at Orchard Road. The record asks learners to bring a one-litre jar for their starter and a container for their loaves. Its allergen information names wheat, rye and sesame. These are synthetic catalogue facts used for learning, rather than a real public offer.

Before choosing any real class, compare the current course page with your availability and confirm details that matter to you with the provider. Do not assume that a class capacity tells you how many spaces remain. A fee also needs its own context; the fictional catalogue does not establish tax treatment or a promotional discount.

For this exercise, the next action is simple: view the BAK-101 details in the demo catalogue and compare the learning outcomes with the question you named at the start. Keep the facts visible while you decide what you want to practise. Learn it. Make it. Take it further.

Classroom draft. Source: data/courses.json, BAK-101 title, level, weeks, fee, when, campus, bring, allergens and learn; brand-voice tagline. No real campaign has been published.
'''
write(5,'assets/article-outline.md','# Article outline\n\n1. Reader need: starter confidence and choosing a structured class.\n2. Connected practice: discuss the BAK-101 outcomes without inventing recipe instructions.\n3. Decision: present the scenario offer accurately and give one course-details CTA.\n\nLength: 500–700 words. No fake testimonial, scarcity or food-safety advice.\n')
write(5,'checkpoints/article.md',article)
for platform in ['instagram','facebook']:
 write(5,f'checkpoints/{platform}.md',f'''# {platform.title()} package — sample classroom draft

Hook: Build a starter routine you understand

Caption: A useful baking course begins with the question you want to answer. In the fictional BAK-101 Artisan Sourdough Bread Baking course, learners practise starter care, shaping and home-oven baking over four weeks. View course details in the demo catalogue.

Hashtags: #Sourdough #BreadBaking #StarterCare #SingaporeBaking #SingaporeFood

Proposed alt text: A sourdough loaf on a cream work surface. No actual photographic asset has been supplied for this caption.

Fact sources: BAK-101.title, weeks and learn. Hook and audience angle are creative proposals. Status: draft; not posted.
''')
write(5,'checkpoints/editorial-review.md','# Editorial review — sample\n\nCheck the hook is fulfilled, offer facts match BAK-101, one CTA appears, voice is practical and alt text refers to the actual visual. Record corrections with the revised file name.\n')
cal='day,platform,pillar,audience,asset,source,cta,owner,dependency,status,hard_sell\n'
for day,pillar,asset,hard in [(1,'teach','content/article.md','no'),(3,'teach','content/instagram.md','no'),(5,'teach','content/facebook.md','no'),(7,'question','content/posts.md','no'),(9,'invite','creative/sourdough/flyer-square.pdf','yes')]:
 source='BAK-101.bring' if day==7 else ('BAK-101.title;weeks;fee;when;campus' if day==9 else 'BAK-101.learn')
 cal+=f'Day {day},Instagram,{pillar},weekend home baker,{asset},{source},View demo details,classroom marketer,{asset},draft,{hard}\n'
write(6,'assets/calendar-template.csv',cal.splitlines()[0]+'\n');write(6,'checkpoints/calendar.csv',cal)
write(6,'assets/content-review-template.csv','post,criterion,evidence,status,correction\nP01,source fidelity,,draft,\nP02,CTA count,,draft,\nP03,hashtags,,draft,\nP04,alt text,,draft,\nP05,promotional balance,,draft,\n')
posts='# Five-post recovery pack — classroom draft\n\n'
for i,angle in enumerate(['Name your starter question','Choose the practice you need','Look at the learning outcomes','What should I bring?','Explore the four-week course'],1):
 posts+=f'## P{i:02d}\n\nHook: {angle}\n\nCaption: '+(['Start with one baking question and use it to choose the learning outcomes that matter to you. The fictional BAK-101 course includes starter care. View the details in the demo catalogue.','A useful class description shows what you can practise. BAK-101 includes mixing, shaping and home-oven baking. View the details in the demo catalogue.','Starter care and dough handling are connected skills in the BAK-101 catalogue. Compare those outcomes with your own learning need. View the details in the demo catalogue.','For the fictional BAK-101 class, the catalogue asks for a one-litre jar and a container for your loaves. View the details in the demo catalogue.','Explore BAK-101 Artisan Sourdough Bread Baking: four weeks, S$680, Saturday 9:30am–12:30pm at Orchard Road in this fictional scenario. View the details in the demo catalogue.'][i-1])+'\n\nHashtags: #Sourdough #BreadBaking #StarterCare #SingaporeBaking #SingaporeFood\n\nProposed alt text: A typographic course-details panel on a cream background. Replace this with the actual image description after selecting the visual.\n\nSource: BAK-101 record. Status: draft.\n\n'
write(6,'checkpoints/posts.md',posts)
contacts='id,email,first_name,consent,interest,suppressed\nC001,learner1@example.com,Ada,yes,bakery,no\nC002,learner2@example.com,Ben,no,bakery,no\nC003,learner3@example.com,Chen,yes,cooking,no\nC004,learner4@example.com,,yes,bakery,no\nC005,learner5@example.com,Dee,yes,bakery,yes\nC006,learner6@example.com,Eli,yes,bakery,no\nC007,learner7@example.com,Fay,,bakery,no\nC008,learner8@example.com,Gus,yes,cooking,yes\n'
write(7,'assets/contacts.csv',contacts)
write(7,'checkpoints/segment.csv','\n'.join([l for l in contacts.splitlines() if l.startswith(('id,','C001,','C004,','C006,'))])+'\n')
sequence='# Three-email sequence — supplied classroom draft, not sent\n\n'
for num,sub,pre,body in [(1,'Build your sourdough routine','Explore a four-week starter-focused class.','The fictional BAK-101 Artisan Sourdough Bread Baking course covers starter care, shaping and home-oven baking. Its catalogue lists four weeks, S$680 and Saturday 9:30am–12:30pm at Orchard Road.'),(2,'What would you like to practise?','Connect your learning question with the course outcomes.','BAK-101 includes feeding a starter, shaping by hand and baking in a home oven. Compare these outcomes with the question you want to answer. The scenario course lasts four weeks.'),(3,'Keep the course details handy','Review the schedule before choosing your next step.','If you are still comparing options, the fictional BAK-101 record lists Saturday 9:30am–12:30pm at Orchard Road, four weeks and S$680. Remaining seats and promotional deadlines are UNKNOWN.')]:
 sequence+=f'## Email {num} — Day {[1,4,8][num-1]} proposed\n\nSubject: {sub}\n\nPreheader: {pre}\n\nHello there,\n\n{body}\n\nCTA: View course details in the demo catalogue.\n\nProduction sender, sender address and unsubscribe configuration: UNKNOWN.\n\n'
write(7,'checkpoints/sequence.md',sequence)
emailtxt='Hello there,\n\nExplore BAK-101 Artisan Sourdough Bread Baking in our fictional classroom catalogue. Practise starter care, shaping and home-oven baking.\n\n4 weeks · S$680 · Saturday 9:30am–12:30pm · Orchard Road\n\nView course details in the demo catalogue: https://example.com/courses/BAK-101?utm_source=email&utm_medium=newsletter&utm_campaign=sourdough_oct_demo&utm_content=e1_a\n\nLearn it. Make it. Take it further.\n\nClassroom draft — not sent. Demo unsubscribe placeholder; provider setup UNKNOWN.\n'
write(7,'checkpoints/email-1.txt',emailtxt)
html='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>BAK-101 email preview — classroom draft</title><body style="margin:0;background:#FBF7F1;color:#2B211B;font:17px/1.6 Arial,sans-serif"><div style="max-width:600px;margin:auto;padding:28px;box-sizing:border-box"><p style="color:#3F6B4A">COOK &amp; BAKE ACADEMY · DEMO</p><h1 style="font-family:Georgia">Build your sourdough routine</h1><p>Hello there,</p><p>Explore BAK-101 Artisan Sourdough Bread Baking in our fictional classroom catalogue. Practise starter care, shaping and home-oven baking.</p><p><strong>4 weeks · S$680</strong><br>Saturday 9:30am–12:30pm<br>Orchard Road</p><p><a style="display:inline-block;background:#3F6B4A;color:#fff;padding:14px;text-decoration:none" href="https://example.com/courses/BAK-101?utm_source=email&amp;utm_medium=newsletter&amp;utm_campaign=sourdough_oct_demo&amp;utm_content=e1_a">View demo course details</a></p><p>Learn it. Make it. Take it further.</p><hr><p style="font-size:14px">Classroom draft — not sent.<br><a href="#demo-unsubscribe">Demo unsubscribe placeholder</a>; provider setup UNKNOWN.</p><p id="demo-unsubscribe" style="font-size:14px">This preview does not unsubscribe a real contact.</p></div></body></html>'''
write(7,'checkpoints/email-1.html',html)
write(7,'checkpoints/send-checklist.md','# Send checklist — classroom draft, NOT SENT\n\nEligibility: consent=yes AND interest=bakery AND suppressed=no. Expected IDs: C001, C004, C006.\n\nProduction sender, sender address, live booking destination and provider unsubscribe configuration: UNKNOWN. Record an owner and next action before production use. Browser preview is not an inbox or deliverability test. No upload, send or schedule is authorised.\n')
results='variant,sent,delivered,unique_clicks,bookings\nA,1000,980,49,7\nB,1000,975,58,8\n'
write(8,'assets/campaign-results.csv',results)
write(8,'assets/metric-definitions.md','# Synthetic metric definitions\n\nDelivery rate = delivered/sent. Unique click rate = unique_clicks/delivered. Click-to-booking conversion = bookings/unique_clicks. Use UNKNOWN when a denominator is zero. The counts are invented teaching data, not academy performance.\n')
write(8,'assets/manifest-template.csv','path,purpose,version,source,status,unresolved_issue\n')
write(8,'checkpoints/metrics.csv','variant,delivery_rate,unique_click_rate,click_to_booking\nA,0.98,0.05,0.142857\nB,0.975,0.059487,0.137931\n')
write(8,'checkpoints/recommendation.md','# Recommendation — supplied classroom draft\n\nObservation: B has 58 unique clicks from 975 deliveries, versus A with 49 from 980. B has a higher observed click rate; A has a slightly higher click-to-booking rate. Limitation: synthetic small counts do not establish statistical confidence or causal uplift. Proposed next test: change only the subject line, retain the audience/body/CTA and review counts over an agreed window. Real performance is UNKNOWN.\n')
write(8,'checkpoints/handoff.md','# Handoff — supplied classroom draft\n\nRetain BAK-101, four weeks, S$680, the Saturday schedule and academy identity across all assets. Record which flyer, article and email files were opened and compared. Add all current paths to manifest.csv. Production sender, live booking link, unsubscribe configuration and numeric launch target remain UNKNOWN. Owner: campaign marketer; next action: confirm with authorised production stakeholders. No message has been sent, scheduled or posted.\n')
write(8,'checkpoints/manifest.csv','path,purpose,version,source,status,unresolved_issue\ncampaign/metrics.csv,synthetic analysis,1.0,campaign-results.csv,classroom draft,real performance UNKNOWN\ncampaign/recommendation.md,next test,1.0,metrics.csv,classroom draft,observation window UNKNOWN\n')
print('Prepared source-only website and eight recovery packs; article words',len(article.split()))

# Proof the actual editable sample text, including motifs and footers.
for n in [3,4]:
 target=folder(n)/'checkpoints';parts=[]
 for ppt in target.glob('flyer-*.pptx'):
  deck=Presentation(ppt);lines=[shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame and shape.text]
  parts.append('## '+ppt.name+'\n\n```text\n'+'\n'.join(lines)+'\n```\n\nSources: BAK-101 title, code, weeks, fee, when and campus. Tagline from brand-voice. Headline, motif and CTA are creative text. No imagery generated.\n')
 (target/('text-proof.md' if n==3 else 'variants-proof.md')).write_text('# Complete visible-text proof — supplied classroom samples\n\n'+'\n'.join(parts))
 if n==3:(target/'flyer-copy.md').write_text('\n'.join(lines)+'\n')

from pathlib import Path
import re,sys,subprocess,json
from docx import Document
from docx.shared import Pt
R=Path(__file__).resolve().parents[1]
for a in json.loads((R/'data/labs.json').read_text()):
 p=R/'labs'/f"lab-{a['num']:02d}-{a['slug']}";doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(11)
 code=False
 for line in (p/'README.md').read_text().splitlines():
  if line.startswith('```'):code=not code;continue
  if not line.strip():continue
  if code:
   para=doc.add_paragraph(line);para.paragraph_format.space_after=Pt(2)
   for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9)
  elif line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));doc.add_heading(line.lstrip('# '),level=min(level,3))
  else:
   plain=re.sub(r'\*\*|`','',line)
   doc.add_paragraph(plain[2:] if plain.startswith('- ') else plain,style='List Bullet' if plain.startswith('- ') else None)
 dest=p/f"C818-Lab-{a['num']:02d}.docx";doc.save(dest)
 r=subprocess.run(['soffice','-env:UserInstallation=file:///tmp/c818-lab-lo','--headless','--convert-to','pdf','--outdir',str(p),str(dest)],capture_output=True,text=True)
 if r.returncode:raise RuntimeError(r.stderr)
 print('Rendered lab',a['num'])
